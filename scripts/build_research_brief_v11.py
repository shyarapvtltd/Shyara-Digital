"""
Planning deliverable v1.1: research-backed developer brief (no code snippets).
Run: python scripts/build_research_brief_v11.py
"""
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = r"e:\Websites\digital-shyara-co-in\Shyara Digital - Studio Editor Developer Brief.docx"


def diagram(doc, title, lines):
    p = doc.add_paragraph()
    r = p.add_run(title + "\n\n" + "\n".join(lines))
    r.font.name = "Consolas"
    r.font.size = Pt(8)


def bullets(doc, items):
    for i in items:
        doc.add_paragraph(i, style="List Bullet")


def numbered(doc, items):
    for n, i in enumerate(items, 1):
        doc.add_paragraph(f"{n}. {i}", style="List Number")


def table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for c, h in enumerate(headers):
        t.rows[0].cells[c].text = h
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            t.rows[ri + 1].cells[ci].text = str(val)
    doc.add_paragraph()


doc = Document()
h = doc.add_heading("Shyara Digital", 0)
h.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_heading("Studio Video Editor — Developer Brief", 1)
doc.add_paragraph("Version 1.1 | Planning research handoff | No code in this document")
doc.add_paragraph(
    "Deliverable of a planning-only phase. Consolidates product decisions, web and open-source "
    "research (Canva animation model, frame-by-frame rendering, stack compatibility), and architecture "
    "for the development team. Canva is a design reference only — we do not use the Canva API. "
    "Implementation is performed by engineering."
)
doc.add_paragraph(
    "Companion files in repo: docs/planning/RESEARCH.md, docs/planning/DEV_TEAM_HANDOFF.md, "
    "docs/planning/FOUNDER_QA_CHECKLIST.md"
)

doc.add_page_break()
doc.add_heading("1. Executive Summary", 1)
doc.add_paragraph(
    "Add a self-serve video invitation platform to digital.shyara.co.in. Customers browse samples, "
    "pay for a template, then open a fullscreen studio where they tap text on the video to personalize "
    "names, dates, and venue. Our servers render the final MP4. MVP is Engagement Sample 1 only."
)
bullets(doc, [
    "Payment is required BEFORE the studio opens — no free editor trial.",
    "Pre-pay preview is YouTube on the Samples page only.",
    "All text animations are built in-house in a shared animation-engine — not imported from Canva.",
    "Admin Dashboard v1 controls templates, prices, coupons — nothing business-critical in React source.",
    "PDF editor is separate (static text, no animations) — not part of MVP studio.",
    "Golden QA: Sample Engagement Invitation 1.mp4 (repository root).",
    "Layer spec: Engagement Sample 1 Timeline and Animation Details.docx.",
])

doc.add_heading("2. Planning Research — Canva Animations (Design Reference)", 1)
doc.add_paragraph(
    "Canva applies preset animations per text element on a video timeline (Element / Text tabs in the "
    "Animate panel). Each animation can run up to 10 seconds. Element timing controls start, duration, "
    "and sequence. There is no supported public API to export Canva’s internal animation curves."
)
doc.add_paragraph("Source: https://www.canva.com/help/animate-designs/")
table(doc, ["Canva preset (design doc)", "Observed behaviour", "Engine slug"], [
    ["Fade + Writing Style Element", "Opacity fade + progressive reveal of elements/letters", "fade_writing_element"],
    ["Skate RTL + Writing Element", "Text glides right-to-left while revealing", "skate_writing_rtl"],
    ["Burst + Writing", "Burst/pop per letter on enter", "burst_writing_enter"],
])
bullets(doc, [
    "Canva speed sliders (15% or 50% in the design spec) are UI percentages — engineering converts to enter.durationMs and exit.durationMs by frame-comparing the reference MP4.",
    "Enter and exit may use different animation types on the same layer.",
    "burst_writing_enter is typically enter-only (date/venue fields).",
    "Respect prefers-reduced-motion in the studio (Canva also reduces motion when OS setting is on).",
    "Typewriter-style presets in Canva reveal one letter at a time — our writing handlers must match grapheme-by-grapheme reveal, not whole-box toggles.",
])

doc.add_heading("3. Frame-by-Frame Rendering Requirements", 1)
doc.add_paragraph(
    "Entry and exit animations reveal text progressively. At every instant t the visible glyphs, opacity, "
    "and transform must match the reference video. Before enter.startMs the layer must be fully invisible "
    "(opacity 0, no ghost pixels). Use grapheme clusters for Hindi/Devanagari and Latin (Intl.Segmenter or equivalent)."
)
numbered(doc, [
    "Implement one TypeScript package (animation-engine) used in browser preview and render worker.",
    "Browser: drive DOM overlay CSS from engine output synced to video.currentTime.",
    "Worker: for frame n at export fps, t = n/fps; run engine; rasterize with node-canvas; pipe RGB24 to FFmpeg stdin.",
    "FFmpeg: H.264 video; copy audio from background MP4 where possible.",
    "Frame index formula (both paths): floor(t_ms * fps / 1000) with integer milliseconds internally.",
    "Empty user fields: omit layer in output.",
    "While user edits one field: freeze all layer animation.",
    "QA: frame grabs at identical timestamps vs Sample Engagement Invitation 1.mp4; founder sign-off (see FOUNDER_QA_CHECKLIST.md).",
])

doc.add_heading("4. Technology Compatibility Matrix (Validated)", 1)
table(doc, ["Component", "Works with", "Constraint / do not use"], [
    ["Vite + React 18 SPA", "Supabase JS, Razorpay Checkout", "No FFmpeg in browser; lazy-load studio route"],
    ["Supabase Auth + Postgres + Storage", "Fastify API, RLS", "Edge Functions cannot spawn FFmpeg"],
    ["Fastify on Render", "BullMQ, Razorpay webhooks", "Verify webhooks on raw request body"],
    ["BullMQ + Redis", "Dedicated worker (Docker)", "Job payload: draftId + quality preset"],
    ["node-canvas + FFmpeg worker", "Same animation-engine package", "Docker: ffmpeg, fontconfig, Noto Devanagari"],
    ["HTML video + DOM overlays", "contenteditable click-to-edit", "Canvas-only editor hurts mobile UX"],
    ["Shared animation-engine (TS)", "DOM + canvas rasterization", "Remotion-only: parity + licensing risk"],
    ["Razorpay INR", "Entitlements table", "Staging test keys separate from prod Supabase"],
    ["pdf-lib PDF path", "Separate product", "Static text only — no animation-engine in MVP"],
])
doc.add_paragraph(
    "Rejected: Supabase Edge + FFmpeg (no shell/binary on hosted Edge). FFmpeg drawtext-only for Hindi "
    "(broken Devanagari without HarfBuzz/Pango). Browser-only export (weak Indic, IP exposure). "
    "Remotion-only primary path (different runtime than DOM studio preview)."
)

doc.add_heading("5. Reference Projects (Patterns Only — Not Dependencies)", 1)
table(doc, ["Project", "URL / note", "Relevance"], [
    ["Remotion", "github.com/remotion-dev/remotion", "Fonts must load before headless frame capture"],
    ["editly", "github.com/mifi/editly", "JSON timeline + FFmpeg + custom canvas layers"],
    ["canvas2video", "github.com/pankod/canvas2video", "Frame stream piped to FFmpeg over background video"],
    ["subtranslate", "github.com/Syrins/subtranslate", "Supabase + API + queue + FFmpeg worker topology"],
    ["node-canvas", "github.com/Automattic/node-canvas", "Pango shaping; ctx.lang for Indic"],
    ["harfbuzzjs", "github.com/harfbuzz/harfbuzzjs", "Fallback if browser vs server glyphs diverge"],
])

doc.add_heading("6. Preview vs Export Parity Risks", 1)
table(doc, ["Risk", "Mitigation"], [
    ["Font metrics differ", "Same TTF/WOFF in @font-face and worker registerFont"],
    ["Time rounding differs", "Integer ms; identical frame index in browser and worker"],
    ["Hindi clusters wrong", "Grapheme segmentation; test frames vs reference MP4"],
    ["1080 proxy vs 4K drift", "Positions in template canvas space; scale at render time"],
    ["Ghost text before enter", "Opacity 0 and zero visible glyphs before enter.startMs"],
    ["RGBA in MP4 pipeline", "Convert canvas to RGB24 before FFmpeg (H.264 has no alpha)"],
])

doc.add_page_break()
doc.add_heading("7. Current Website vs Target", 1)
table(doc, ["Area", "Today", "Target"], [
    ["Frontend", "Vite + React 18 SPA", "Same + /studio + /admin routes"],
    ["Samples", "Hardcoded YouTube IDs", "API-driven template cards"],
    ["Auth", "None", "Supabase email + Google"],
    ["Payments", "WhatsApp only", "Razorpay INR"],
    ["Database", "None", "Supabase Postgres"],
    ["Video editor", "None", "Click-to-edit fullscreen studio"],
    ["Export", "None", "Server FFmpeg render 60fps"],
])

doc.add_heading("8. Complete User Journey", 1)
numbered(doc, [
    "User browses Samples and watches YouTube marketing preview.",
    "User clicks Customize on a template.",
    "User signs up or logs in; must accept Terms at registration.",
    "User pays through Razorpay; optional influencer coupon (hidden checkout field, Admin-created codes only).",
    "System records entitlement (purchase + download quota, default 3).",
    "User opens /studio/:templateId — fullscreen 9:16 video, no watermark.",
    "User plays (loops), pauses, scrubs; when paused, faint boxes show editable regions.",
    "User clicks text — pause — inline edit; preview live; all motion frozen while one field active.",
    "User Done or tap outside; Export with quality choice and estimated file size.",
    "Server renders if needed; progress in UI + email; login-gated download (each download uses 1 quota).",
])
diagram(doc, "Figure 1 — User flow", [
    "  [Samples + YouTube] -> [Sign up] -> [Pay] -> [Studio editor] -> [Export MP4]",
])

doc.add_heading("9. Studio Editor — Product Requirements", 1)
doc.add_heading("9.1 Layout", 2)
bullets(doc, [
    "Almost entire screen is 9:16 video (letterboxed).",
    "Minimal chrome: Play, Pause, timeline scrubber, Done, Export.",
    "No sidebar field list as primary UI; Tab through layers required (WCAG 2.1 AA).",
])
doc.add_heading("9.2 Click-to-edit behaviour", 2)
bullets(doc, [
    "Click text → auto-pause → contenteditable inline on that layer.",
    "Faint boxes on all editable layers when paused; brand-color border when selected.",
    "Deselect: Done + tap outside.",
    "Freeze all animation while editing one field; loop video; free scrub.",
    "Real-time overlay update; mobile: video scrolls above keyboard; exact hit targets only.",
    "Overlap: top z-index; multiple hits → picker menu.",
    "Warn on navigate away; auto-save draft after payment.",
])
doc.add_heading("9.3 Engagement Sample 1 rules", 2)
table(doc, ["Page", "Time (s)", "Content", "Animations (design doc)"], [
    ["1", "0–6", "Baked in video — NOT editable", "N/A"],
    ["2", "6–12", "Partner 1 + Partner 2 names", "Fade + Writing Element, 15%, enter+exit; Partner 2 same window as Partner 1"],
    ["3", "12–18", "Large names + parent names", "Names: Fade 15%; Parents: Skate RTL 50%, enter+exit"],
    ["4", "18–24", "Month, Day, Date, Time, Year, Venue", "Burst + Writing, enter-focused"],
])
bullets(doc, [
    "Example positions (transcribe all 16 from Timeline docx): Partner 1 Page 2 — 372.3×120.3 @ (353.8, 444.5); Partner 2 — 292.8×118.4 @ (393.6, 629.2).",
    "Fonts: Lucien Schoenschrift CAT, Lora, Futura per design doc.",
])
doc.add_heading("9.4 User restrictions", 2)
bullets(doc, [
    "Text content only — not layout, fonts, colors, or animation.",
    "No emojis; hard character limits; empty fields hidden in output.",
    "No photos in self-serve editor.",
])

doc.add_heading("10. Technology Stack", 1)
table(doc, ["Layer", "Technology", "Purpose"], [
    ["Frontend", "Vite, React 18, TypeScript", "Marketing + Studio + Admin"],
    ["UI", "Tailwind, shadcn/ui", "Components"],
    ["Routing", "React Router", "Navigation"],
    ["Auth and DB", "Supabase", "Users, data, files, RLS"],
    ["API", "Node.js + Fastify", "Business logic, webhooks"],
    ["Queue", "BullMQ + Redis", "Render jobs"],
    ["Worker", "Node + FFmpeg + node-canvas", "MP4 export"],
    ["Animation", "TypeScript animation-engine", "Preview + export parity"],
    ["Payments", "Razorpay", "India INR"],
    ["Email", "Resend or SendGrid", "Notifications"],
    ["Analytics", "GA4", "Funnel"],
    ["Host", "Render", "Static + API + worker"],
])

doc.add_heading("11. System Architecture", 1)
diagram(doc, "Figure 2 — System architecture", [
    " +-------------+     +-------------+     +-------------+",
    " |   Browser   |     | API Fastify |     |   Worker    |",
    " | Studio DOM  |<--->| entitlements|<--->| canvas+FFmpeg|",
    " | + engine    |     | drafts/jobs |     | + engine    |",
    " +------+------+     +------+------+     +------+------+",
    "        +-------------------+-------------------+",
    "                            v",
    "                 +---------------------+",
    "                 | Supabase DB+Storage|",
    "                 +---------------------+",
])
bullets(doc, [
    "Editor loads 1080p proxy; export can use 4K master when selected.",
    "v1: manual upload of both 4K and 1080p per template.",
    "Downloads require login — no public shareable URLs.",
])

doc.add_page_break()
doc.add_heading("12. Template Data (definition_json)", 1)
doc.add_paragraph("Each template version is one JSONB document. Studio and worker read only this JSON.")
bullets(doc, [
    "canvasWidth, canvasHeight, durationMs, fps (60 export), bakedIntroEndMs (6000 for Sample 1).",
    "backgroundVideoUrl1080, backgroundVideoUrl4K.",
    "textLayers: x, y, width, height, fontFamily, fontSize, fontWeight, color, maxChars, maxLines, label, zIndex.",
    "Per layer: enter { type, startMs, durationMs, params }, static { startMs, endMs }, exit { type, startMs, durationMs, params }.",
    "v1 types only: fade_writing_element, skate_writing_rtl, burst_writing_enter.",
    "Launch blocker: all three handlers must match reference MP4 before catalog expansion.",
])

doc.add_heading("13. Text Animations — Built In-House (CRITICAL)", 1)
table(doc, ["Phase", "Meaning", "Fields"], [
    ["Enter", "Text appears", "enter.type, enter.startMs, enter.durationMs, enter.params"],
    ["Static", "Fully visible", "static.startMs, static.endMs"],
    ["Exit", "Text leaves", "exit.type, exit.startMs, exit.durationMs"],
])
doc.add_paragraph(
    "Dev team implements handlers: (layer config, t_ms) → opacity, translate, scale, visibleGraphemeCount. "
    "Browser applies as CSS; worker draws on canvas. See DEV_TEAM_HANDOFF.md for suggested build order."
)

doc.add_heading("14. Commerce and Policies", 1)
bullets(doc, [
    "India INR Razorpay only at launch.",
    "Pay before studio — route blocked without entitlement.",
    "Default 3 downloads per template (Admin override).",
    "Every download costs 1 quota even if cached.",
    "Failed render: no quota; WhatsApp support with order/job context; manual delivery within 48h.",
    "Refund stated at checkout; no refund after first export.",
    "Draft retention 30 days; rendered files 30 days; self-serve account deletion.",
])

doc.add_heading("15. Admin Dashboard (v1 required)", 1)
bullets(doc, [
    "Categories, templates, prices, thumbnails, YouTube preview IDs.",
    "Template versions, JSON import/export, text layers, animation timings.",
    "Font uploads, export quality presets, coupons, orders, render jobs, manual fulfillment.",
    "Platform settings; admin role on /admin/*.",
])

doc.add_heading("16. Environments", 1)
table(doc, ["", "Staging", "Production"], [
    ["Supabase", "Separate project", "Separate project"],
    ["Razorpay", "Test keys only", "Live keys"],
    ["Rule", "Full E2E QA first", "Never mix credentials"],
])

doc.add_heading("17. Suggested Dev Team Build Order (Reference)", 1)
numbered(doc, [
    "Supabase schema + staging + RLS.",
    "animation-engine + golden-frame unit tests for three v1 types.",
    "Engagement Sample 1 definition_json + fonts + 4K/1080p assets.",
    "Studio click-to-edit UI wired to engine.",
    "Razorpay pay gate and entitlements.",
    "Render worker + quality presets + login-gated download.",
    "Admin dashboard v1.",
    "Founder QA per FOUNDER_QA_CHECKLIST.md.",
])

doc.add_heading("18. Team Skills", 1)
table(doc, ["Role", "Focus"], [
    ["Frontend", "React, video API, DOM overlays, contenteditable, Supabase, Razorpay, WCAG"],
    ["Backend", "Fastify, Postgres, webhooks, queues, secure downloads"],
    ["Media worker", "FFmpeg, 60fps pipeline, node-canvas, Indic fonts"],
    ["Shared", "animation-engine, JSON schema, golden-frame QA"],
])

doc.add_heading("19. Out of Scope (MVP)", 1)
bullets(doc, [
    "Canva API; general video editor; user photo upload; drag/move/resize text.",
    "Animated PDF editor; international payments; B2B bulk licensing.",
])

doc.add_heading("20. Planning Phase Boundary", 1)
bullets(doc, [
    "This document and companion planning files are the output of planning-only work.",
    "No application code, migrations, or deployments were produced in the planning phase.",
    "Implementation tracking belongs to the development team.",
])

doc.add_paragraph("End of document.")
doc.save(OUT)
print("Saved:", OUT)
