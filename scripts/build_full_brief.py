from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
OUT = r"e:\Websites\digital-shyara-co-in\Shyara Digital - Studio Editor Developer Brief.docx"

def diagram(doc, title, lines):
    p = doc.add_paragraph()
    r = p.add_run(title + chr(10) + chr(10).join(lines))
    r.font.name = "Consolas"
    r.font.size = Pt(8)

def bullets(doc, items):
    for i in items:
        doc.add_paragraph(i, style="List Bullet")

def numbered(doc, items):
    for n, i in enumerate(items, 1):
        doc.add_paragraph(str(n) + ". " + i, style="List Number")

def table(doc, headers, rows):
    t = doc.add_table(rows=1+len(rows), cols=len(headers))
    t.style = "Table Grid"
    for c, h in enumerate(headers):
        t.rows[0].cells[c].text = h
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            t.rows[ri+1].cells[ci].text = str(val)
    doc.add_paragraph()

doc = Document()
h = doc.add_heading("Shyara Digital", 0)
h.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_heading("Studio Video Editor - Developer Brief", 1)
doc.add_paragraph("Version 1.0 | Internal development team | No code in this document")
doc.add_paragraph(
    "This brief explains what to build, required behaviour, technology choices, and how text animations "
    "must be created in-house. Canva is only a design reference - we do not use Canva API."
)
doc.add_page_break()

doc.add_heading("1. Executive Summary", 1)
doc.add_paragraph(
    "Add a self-serve video invitation platform to digital.shyara.co.in. Customers browse samples, "
    "pay for a template, then open a fullscreen studio where they tap text on the video to personalize "
    "names, dates, and venue. Our servers render the final MP4. MVP is Engagement Sample 1 only."
)
bullets(doc, [
    "Payment is required BEFORE the studio opens - no free editor trial.",
    "Pre-pay preview is YouTube on the Samples page only.",
    "All text animations are built by our team in a shared animation-engine - not imported from Canva.",
    "Admin Dashboard v1 controls templates, prices, coupons - nothing business-critical in React source code.",
    "PDF editor is separate (static text, no animations) - not part of MVP studio.",
])

doc.add_heading("2. Current Website vs Target", 1)
table(doc, ["Area", "Today", "Target"], [
    ["Frontend", "Vite + React 18 SPA", "Same + /studio + /admin routes"],
    ["Samples", "Hardcoded YouTube IDs", "API-driven template cards"],
    ["Auth", "None", "Supabase email + Google"],
    ["Payments", "WhatsApp only", "Razorpay INR"],
    ["Database", "None", "Supabase Postgres"],
    ["Video editor", "None", "Click-to-edit fullscreen studio"],
    ["Export", "None", "Server FFmpeg render 60fps"],
])

doc.add_heading("3. Complete User Journey", 1)
numbered(doc, [
    "User browses Samples and watches YouTube marketing preview.",
    "User clicks Customize on a template.",
    "User signs up or logs in. Must accept Terms at registration.",
    "User pays through Razorpay checkout. Optional influencer coupon in hidden checkout field.",
    "System records entitlement (purchase + download quota, default 3).",
    "User opens /studio/:templateId - fullscreen 9:16 video, no watermark.",
    "User plays video (loops), pauses, scrubs timeline freely.",
    "When paused, faint boxes show editable text regions.",
    "User clicks a text region - video pauses - inline typing on video.",
    "Preview updates live. All motion frozen while one field is active.",
    "User presses Done or taps outside to finish field.",
    "User clicks Export, selects 4K / 2K / 1080p / WhatsApp-optimized, sees estimated file size.",
    "Server renders if needed; progress in UI + email when ready.",
    "User downloads while logged in. Each download uses 1 quota (even if cached).",
])
diagram(doc, "Figure 1 - User flow", [
    "  [Samples + YouTube] -> [Sign up] -> [Pay] -> [Studio editor] -> [Export MP4]",
])

doc.add_page_break()
doc.add_heading("4. Studio Editor - Product Requirements", 1)
doc.add_heading("4.1 Layout", 2)
bullets(doc, [
    "Almost entire screen is the video (9:16 letterboxed).",
    "Minimal chrome: Play, Pause, timeline scrubber, Done, Export.",
    "No sidebar form with field list as primary UI (Tab focus on layers still required for accessibility).",
])
doc.add_heading("4.2 Click-to-edit behaviour", 2)
bullets(doc, [
    "Click text -> auto-pause -> contenteditable inline on that text layer.",
    "Faint boxes around all editable layers when paused.",
    "Selected layer: solid brand-color border.",
    "Deselect: Done button + tap outside.",
    "Freeze all layer animation while one field is being edited.",
    "Loop full video; scrub any time.",
    "Real-time overlay update as user types.",
    "Mobile: scroll video above keyboard; tap target = exact box only.",
    "Tab through layers with focus ring (WCAG 2.1 AA).",
    "Warn on navigate away; auto-save draft after payment.",
])
doc.add_heading("4.3 Engagement Sample 1 rules", 2)
bullets(doc, [
    "Page 1 (0-6s): baked into background MP4 - NOT editable. No hit targets 0-6s.",
    "Page 2 (6-12s): Partner 1 + Partner 2 names (Partner 2 same duration as Partner 1).",
    "Page 3 (12-18s): Large names + parent names.",
    "Page 4 (18-24s): Month, Day, Date, Time, Year, Venue.",
    "Golden QA file: Sample Engagement Invitation 1.mp4 (repo root).",
    "Spec: Engagement Sample 1 Timeline and Animation Details.docx",
])
doc.add_heading("4.4 User restrictions", 2)
bullets(doc, [
    "Text content only - not layout or animation.",
    "No emojis. Hard character limits. Empty fields hidden in output.",
    "No photos in self-serve editor.",
])

doc.add_heading("5. Technology Stack", 1)
table(doc, ["Layer", "Technology", "Purpose"], [
    ["Frontend", "Vite, React 18, TypeScript", "Marketing + Studio + Admin"],
    ["UI", "Tailwind, shadcn/ui", "Components"],
    ["Routing", "React Router", "Page navigation"],
    ["Auth and DB", "Supabase", "Users, data, files, RLS"],
    ["API", "Node.js + Fastify", "Business logic, webhooks"],
    ["Queue", "BullMQ + Redis", "Render jobs"],
    ["Worker", "Node + FFmpeg + node-canvas", "MP4 export"],
    ["Animation", "TypeScript animation-engine package", "Preview + export parity"],
    ["Payments", "Razorpay", "India INR"],
    ["Email", "Resend or SendGrid", "Notifications"],
    ["Analytics", "GA4", "Funnel"],
    ["Host", "Render", "Static + API + worker"],
])

doc.add_heading("6. System Architecture", 1)
doc.add_paragraph(
    "Recommended: HTML video element + DOM text overlays for the studio. "
    "Shared animation-engine drives CSS transforms in browser. "
    "Server uses same engine + canvas per frame + FFmpeg for final file."
)
diagram(doc, "Figure 2 - System architecture", [
    " +-------------+     +-------------+     +-------------+",
    " |   Browser   |     | API Fastify |     |   Worker    |",
    " | Studio DOM  |<--->| entitlements|<--->| canvas+FFmpeg|",
    " | + engine    |     | drafts/jobs |     | + engine    |",
    " +------+------+     +------+------+     +------+------+",
    "        |                   |                   |",
    "        +-------------------+-------------------+",
    "                            v",
    "                 +---------------------+",
    "                 | Supabase DB+Storage|",
    "                 +---------------------+",
])
diagram(doc, "Figure 3 - Studio browser stack", [
    " +----------------------------------+",
    " | Play | Pause | ===time=== | Export |",
    " +----------------------------------+",
    " |   [ Text overlay layers DOM ]    |  <- click + edit",
    " |   [ HTML video 1080p proxy ]     |",
    " +----------------------------------+",
])
bullets(doc, [
    "Editor loads 1080p proxy for speed.",
    "Export uses 4K master when user picks 4K.",
    "Dev uploads BOTH 4K master and 1080p proxy per template (v1).",
    "Downloads require login - no public shareable URLs.",
])

doc.add_page_break()
doc.add_heading("7. Template Data (definition_json)", 1)
doc.add_paragraph("Each template version is one JSON document in the database. Studio and worker read only this JSON.")
bullets(doc, [
    "canvasWidth, canvasHeight (per template, 9:16)",
    "durationMs, fps (60 for export)",
    "backgroundVideoUrl1080, backgroundVideoUrl4K",
    "bakedIntroEndMs (e.g. 6000 for Engagement Sample 1)",
    "fonts: Lucien Schoenschrift CAT, Lora, Futura (+ bold via fontWeight)",
    "textLayers: x, y, width, height, font, size, color, maxChars, maxLines, label, zIndex",
    "per-layer animation with separate enter and exit types and timings in milliseconds",
])
doc.add_paragraph("Position values come from Canva screenshots in the spec doc (Width, Height, X, Y in pixels). Engineering transcribes into JSON.")

doc.add_heading("8. Text Animations - Built In-House (CRITICAL)", 1)
doc.add_paragraph(
    "We do NOT use Canva API or runtime Canva export. Designers use Canva as reference. "
    "Developers implement every animation in animation-engine. Same package runs in browser preview AND server export."
)
doc.add_heading("8.1 Per-layer time model", 2)
table(doc, ["Phase", "Meaning", "Fields"], [
    ["Enter", "Text appears", "enter.type, enter.startMs, enter.durationMs, enter.params"],
    ["Static", "Fully visible", "static.startMs, static.endMs"],
    ["Exit", "Text leaves", "exit.type, exit.startMs, exit.durationMs (can differ from enter)"],
])
bullets(doc, [
    "Layer invisible before enter.startMs.",
    "Each layer has its own enter.startMs (manual stagger).",
    "Convert Canva slide-bar % (15% or 50%) to milliseconds using reference MP4.",
])

doc.add_heading("8.2 Three animation types - Version 1", 2)
table(doc, ["Engine slug", "Canva name", "Exit?", "Used on"], [
    ["fade_writing_element", "Fade + Writing Style Element", "Yes", "Names, titles"],
    ["skate_writing_rtl", "Skate RTL + Writing Element", "Yes", "Parent names"],
    ["burst_writing_enter", "Burst + Writing", "Enter only", "Date and venue fields"],
])
doc.add_paragraph("Launch blocker: all three must match reference MP4 before release. More types added later.")

doc.add_heading("8.3 How to implement animations (step by step)", 2)
numbered(doc, [
    "Create packages/animation-engine (TypeScript) in monorepo.",
    "For each animation slug, write a handler: input time t + layer config -> opacity, position, scale.",
    "fade_writing_element: fade + character/element writing reveal over enter.durationMs; reverse on exit.",
    "skate_writing_rtl: horizontal slide right-to-left with fade; 50% speed in doc.",
    "burst_writing_enter: burst/pop on enter only; no exit animation.",
    "Browser: apply engine output as CSS on DOM layers synced to video.currentTime.",
    "Server: for each frame at 60fps, run engine, rasterize text on canvas, feed frames to FFmpeg.",
    "FFmpeg merges frames with background video; copy audio from background MP4.",
    "Unit tests at fixed timestamps.",
    "Visual QA: compare frames to Sample Engagement Invitation 1.mp4.",
    "Founder sign-off required.",
])

doc.add_heading("8.4 Preview vs export", 2)
table(doc, ["", "Browser", "Server"], [
    ["Video", "1080p proxy", "Up to 4K"],
    ["Text", "DOM + CSS", "Canvas frames"],
    ["Math", "animation-engine", "Same animation-engine"],
    ["Output", "Preview only", "Final H.264 MP4"],
])

doc.add_page_break()
doc.add_heading("9. Commerce and Policies", 1)
bullets(doc, [
    "India INR Razorpay only at launch.",
    "Pay before studio - route blocked without entitlement.",
    "Default 3 downloads per template (Admin can override).",
    "Every download costs 1 quota even if file already rendered.",
    "Failed render: no quota charge; WhatsApp support with order/job context; manual delivery within 48h.",
    "Coupons: Admin-created influencer codes; hidden at checkout.",
    "Refund: stated at checkout; no refund after first export; manual support otherwise.",
    "Self-serve account deletion.",
    "Draft retention 30 days; rendered files kept 30 days on server.",
    "Repurchase = new entitlement and new quota.",
])

doc.add_heading("10. Admin Dashboard (v1 required)", 1)
bullets(doc, [
    "Categories, templates, prices, thumbnails, YouTube preview IDs",
    "Template versions and JSON import/export",
    "Text layers and animation timings (numeric)",
    "Font uploads",
    "Export quality presets",
    "Coupons and orders",
    "Render job management and manual fulfillment upload",
    "Platform settings",
    "Admin role on /admin/* routes",
])

doc.add_heading("11. Environments", 1)
table(doc, ["", "Staging", "Production"], [
    ["Supabase", "Separate project", "Separate project"],
    ["Razorpay", "Test keys only", "Live keys"],
    ["Rule", "Full E2E QA first", "Never mix credentials"],
])

doc.add_heading("12. MVP Build Order", 1)
numbered(doc, [
    "Supabase + staging + API skeleton",
    "animation-engine with 3 types + tests",
    "Engagement Sample 1 definition_json + font/video assets",
    "Studio click-to-edit UI",
    "Razorpay gate before studio",
    "Render worker + quality presets + download",
    "Admin dashboard",
    "Founder QA vs reference MP4",
])

doc.add_heading("13. Team Skills", 1)
table(doc, ["Role", "Focus"], [
    ["Frontend", "React, video API, DOM overlays, contenteditable, Supabase, Razorpay, WCAG"],
    ["Backend", "Fastify, Postgres, webhooks, queues, secure downloads"],
    ["Media worker", "FFmpeg, 60fps pipeline, canvas, Indic fonts"],
    ["Shared", "animation-engine, JSON schema, golden-frame QA"],
])

doc.add_heading("14. Out of Scope for MVP", 1)
bullets(doc, [
    "General purpose video editor",
    "Canva API",
    "User photo upload in studio",
    "Drag/move/resize text",
    "PDF click-to-edit",
    "International payments",
    "B2B bulk licensing",
])

doc.add_paragraph("End of document.")
doc.save(OUT)
print("Saved comprehensive brief:", OUT)
