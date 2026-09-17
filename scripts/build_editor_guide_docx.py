"""Generate Shyara Digital - Studio Editor Implementation Guide.docx from markdown source."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path(r"e:\Websites\digital-shyara-co-in")
MD = ROOT / "docs" / "planning" / "EDITOR_IMPLEMENTATION_GUIDE.md"
OUT = ROOT / "Shyara Digital - Studio Editor Implementation Guide.docx"


def add_from_md(doc: Document, text: str) -> None:
    """Simple markdown-to-docx: headings, bullets, tables, paragraphs."""
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), 0)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), 1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), 2)
        elif line.startswith("#### "):
            doc.add_heading(line[5:].strip(), 3)
        elif line.startswith("| ") and i + 1 < len(lines) and lines[i + 1].startswith("|-"):
            headers = [c.strip() for c in line.strip("|").split("|")]
            i += 2
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append([c.strip() for c in lines[i].strip("|").split("|")])
                i += 1
            t = doc.add_table(rows=1 + len(rows), cols=len(headers))
            t.style = "Table Grid"
            for c, h in enumerate(headers):
                t.rows[0].cells[c].text = h
            for ri, row in enumerate(rows):
                for ci, val in enumerate(row):
                    if ci < len(headers):
                        t.rows[ri + 1].cells[ci].text = val
            doc.add_paragraph()
            continue
        elif line.startswith("- [ ] "):
            doc.add_paragraph(line[6:].strip(), style="List Bullet")
        elif line.startswith("- "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        elif line.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(block))
            r.font.name = "Consolas"
            r.font.size = Pt(8)
        elif line.strip() == "---":
            doc.add_page_break()
        elif line.strip():
            doc.add_paragraph(line.strip())
        i += 1


def main() -> None:
    md = MD.read_text(encoding="utf-8")
    # Skip title block — add branded cover
    doc = Document()
    h = doc.add_heading("Shyara Digital", 0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading("Studio Video Editor — Implementation Guide", 1)
    doc.add_paragraph(
        "For the development team | How to build the editor, animations, and render pipeline | Version 1.0"
    )
    doc.add_paragraph(
        "This is the primary technical document for building the click-to-edit studio. "
        "Canva is design reference only. Golden QA: Sample Engagement Invitation 1.mp4."
    )
    doc.add_page_break()
    # Body starts after first ---
    body = md.split("---", 1)[-1] if "---" in md else md
    add_from_md(doc, body)
    doc.save(OUT)
    print("Saved:", OUT)


if __name__ == "__main__":
    main()
