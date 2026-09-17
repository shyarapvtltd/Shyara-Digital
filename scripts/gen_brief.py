from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
OUTPUT = r"e:\Websites\digital-shyara-co-in\Shyara Digital - Studio Editor Developer Brief.docx"
def add_diagram_box(doc, title, lines):
    p = doc.add_paragraph()
    r = p.add_run(title + "\n" + "\n".join(lines))
    r.font.name = "Consolas"
    r.font.size = Pt(9)
def build():
    doc = Document()
    t = doc.add_heading("Shyara Digital — Video Studio & Platform", 0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Developer Brief — Requirements, Technology & Architecture (v1.0)")
    doc.add_paragraph("For internal development team. No implementation code in this document.")
    doc.add_page_break()
    doc.add_heading("1. Executive Summary", 1)
    doc.add_paragraph("Build a self-serve invitation video platform on the existing marketing site. Customers browse samples, pay, then edit text by tapping directly on a fullscreen video. The system renders a final MP4 (4K/2K/1080p/WhatsApp) on our servers. MVP template: Engagement Sample 1.")
    for x in ["MVP is video studio only first.", "Animations are built in-house — not Canva API.", "Admin Dashboard v1 — no hardcoded prices or templates.", "Payment before editor — no free studio trial."]:
        doc.add_paragraph(x, style="List Bullet")
    doc.add_heading("2. User Journey", 1)
    for i,s in enumerate(["Browse Samples (YouTube preview).","Click Customize.","Sign up / login + accept terms.","Pay via Razorpay (coupon optional at checkout).","Studio opens — fullscreen 9:16 video.","Play, pause, scrub timeline.","Click text → pause → edit inline.","Export → pick quality → render/download.","Each download uses 1 from quota (default 3)."],1):
        doc.add_paragraph(f"{i}. {s}", style="List Number")
    add_diagram_box(doc, "Flow:", ["  Samples → Sign up → Pay → Studio → Export/Download"])
    doc.add_heading("3. Studio Editor Requirements", 1)
    for x in ["Video-first fullscreen — not a side form.","Click text to edit; auto-pause on click.","Faint boxes when paused; brand border when selected.","Done button + tap outside to deselect.","Freeze all motion while editing one field.","Loop video; free scrub.","Real-time text preview.","Page 1 of Engagement Sample 1 (0-6s): baked in video — not editable.","Mobile: scroll video above keyboard.","Tab through fields for accessibility (WCAG 2.1 AA).","Warn on leave; auto-save draft after payment."]:
        doc.add_paragraph(x, style="List Bullet")
    doc.add_heading("4. Technology Stack", 1)
    rows=[("Frontend","Vite, React 18, TypeScript, Tailwind, shadcn/ui"),("Auth/DB","Supabase Auth, Postgres, Storage, RLS"),("API","Node.js Fastify on Render"),("Queue","BullMQ + Redis (Upstash)"),("Worker","Node + FFmpeg + node-canvas"),("Payments","Razorpay"),("Email","Resend or SendGrid"),("Analytics","GA4"),("Deploy","Render static + API + worker")]
    tbl=doc.add_table(rows=1+len(rows),cols=2); tbl.style="Table Grid"
    tbl.rows[0].cells[0].text="Layer"; tbl.rows[0].cells[1].text="Technology"
    for i,(a,b) in enumerate(rows): tbl.rows[i+1].cells[0].text=a; tbl.rows[i+1].cells[1].text=b
    doc.add_heading("5. System Architecture", 1)
    doc.add_paragraph("Browser: HTML video + DOM text overlays + shared animation-engine for preview. Server: same animation-engine + canvas per frame + FFmpeg for final MP4. API: entitlements, drafts, orders, render queue, authenticated downloads.")
    add_diagram_box(doc, "Architecture:", ["  [Browser Studio] <-> [Fastify API] <-> [Supabase]","                        |","                        v","                 [Render Worker + FFmpeg]"])
    doc.add_heading("6. Template Data (definition_json)", 1)
    for x in ["Each template version stored as JSON in database.","Canvas width/height per template (9:16).","Background: 1080p for editor, 4K for export (dev uploads both).","Text layers: position, font, size, color, max chars, labels.","Per-layer animation: enter/static/exit with independent enter and exit types.","bakedIntroEndMs for non-editable intro (6000ms for Engagement Sample 1)."]:
        doc.add_paragraph(x, style="List Bullet")
    doc.add_heading("7. Animations — Built In-House (Critical)", 1)
    doc.add_paragraph("We do NOT use Canva at runtime. Design uses Canva as reference only. Engineering implements a shared animation-engine package used in browser preview AND server export.")
    doc.add_heading("7.1 Three animation types for v1", 2)
    tbl2=doc.add_table(rows=4,cols=4); tbl2.style="Table Grid"
    for j,h in enumerate(["Type","Canva name","Exit?","Used for"]): tbl2.rows[0].cells[j].text=h
    data=[("fade_writing_element","Fade + Writing Element","Yes","Names, titles — 15-50% speed"),("skate_writing_rtl","Skate RTL + Writing","Yes","Parent names"),("burst_writing_enter","Burst + Writing","Enter only","Month, day, date, time, year, venue")]
    for i,row in enumerate(data,1):
        for j,v in enumerate(row): tbl2.rows[i].cells[j].text=v
    doc.add_heading("7.2 How animations work", 2)
    for x in ["Each layer has enter (type, startMs, durationMs), static (startMs, endMs), exit (optional separate type).","Before enter: layer invisible.","Engine returns opacity/transform at time t.","Browser applies via CSS on DOM layers synced to video.currentTime.","Server draws each frame at 60fps using canvas + same engine math.","Calibrate Canva slide-bar % to milliseconds using Sample Engagement Invitation 1.mp4.","Unit tests + founder visual sign-off required before launch."]:
        doc.add_paragraph(x, style="List Bullet")
    doc.add_heading("8. Commerce & Admin", 1)
    for x in ["India INR Razorpay only.","Pay before studio — no editor without entitlement.","Admin: templates, JSON, fonts, coupons, orders, render jobs, settings.","Coupons: admin-only, hidden at checkout.","Failed render: WhatsApp support, manual file within 48h.","Separate staging Supabase + Razorpay test."]:
        doc.add_paragraph(x, style="List Bullet")
    doc.add_heading("9. MVP Build Order", 1)
    for i,s in enumerate(["Foundation: Supabase, API, animation-engine (3 types).","Engagement Sample 1 JSON + assets.","Studio click-to-edit UI.","Razorpay gate.","Render worker + export qualities.","Admin dashboard.","Founder QA vs reference MP4."],1):
        doc.add_paragraph(f"{i}. {s}", style="List Number")
    doc.add_heading("10. Reference Files", 1)
    for x in ["Sample Engagement Invitation 1.mp4 — golden QA reference","Engagement Sample 1 Timeline and Animation Details.docx — specs and position screenshots"]:
        doc.add_paragraph(x, style="List Bullet")
    doc.save(OUTPUT)
    print("Saved", OUTPUT)
build()
